from flask import Flask, render_template, request, send_file
import os
from docx import Document
from io import BytesIO

app = Flask(__name__)

# Rota principal que carrega o formulário
@app.route('/')
def index():
    return render_template('index.html')

# Rota para processar o formulário e gerar o documento
@app.route('/gerar-documento', methods=['POST'])
def gerar_documento():
    # Capturar os dados do formulário
    nome_completo = request.form['nomecompleto']
    rg = request.form['rg']
    cpf = request.form['cpf']
    unidade = request.form['unidade']
    modulo = request.form['modulo']
    responsavel = request.form['responsavel']
    carga = request.form['carga']
    data = request.form['data']
    resumo = request.form['resumo']
    sintome = request.form['sintome']

    # Caminho para o modelo de documento
    caminho_modelo = os.path.join('documentos', 'TERMO_DE_VALIDAÇÃO_E_PRESENÇA_NO_MÓDULO_DE_TREINAMENTO.docx')

    # Abrir o documento modelo
    doc = Document(caminho_modelo)

    # Substituir os placeholders no documento
    for paragrafo in doc.paragraphs:
        paragrafo.text = paragrafo.text.replace('[NOMECOMPLETO]', nome_completo)
        paragrafo.text = paragrafo.text.replace('[RG]', rg)
        paragrafo.text = paragrafo.text.replace('[CPF]', cpf)
        paragrafo.text = paragrafo.text.replace('[UNIDADE/TERRITORIO]', unidade)
        paragrafo.text = paragrafo.text.replace('[MODULO]', modulo)
        paragrafo.text = paragrafo.text.replace('[RESPONSAVEL]', responsavel)
        paragrafo.text = paragrafo.text.replace('[CARGA]', carga)
        paragrafo.text = paragrafo.text.replace('[DATA]', data)
        paragrafo.text = paragrafo.text.replace('[RESUMO]', resumo)
        paragrafo.text = paragrafo.text.replace('[SINTOME]', sintome)

    # Criar um arquivo temporário para salvar o documento preenchido
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    # Enviar o arquivo preenchido para o download
    return send_file(buffer, as_attachment=True, download_name='termo_preenchido.docx', mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

# Apenas uma vez o app.run
if __name__ == '__main__':
    app.run(debug=True)
